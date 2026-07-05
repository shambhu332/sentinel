"""Split SENTINEL_Coursework.docx into two files for Criteria 1 and Criteria 4."""
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "SENTINEL_Coursework.docx"


def collect_body_items(doc):
    """Return list of (kind, element, para_index_if_p) walking body in order."""
    items = []
    p_idx = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            items.append(("p", child, p_idx))
            p_idx += 1
        elif tag == "tbl":
            items.append(("tbl", child, None))
        else:
            items.append((tag, child, None))
    return items


def build(out_path, title_lines, keep_p_range):
    """Build a new docx keeping only body items whose paragraph index falls in
    keep_p_range; surrounding tables that sit between kept paragraphs are kept too."""
    src = Document(SRC)
    items = collect_body_items(src)
    lo, hi = keep_p_range

    # Identify kept paragraph positions in the items list
    kept_positions = [i for i, (k, _, idx) in enumerate(items)
                      if k == "p" and idx is not None and lo <= idx <= hi]
    if not kept_positions:
        raise SystemExit("nothing kept")
    first_kept, last_kept = kept_positions[0], kept_positions[-1]

    # Identify items to remove: everything outside [first_kept, last_kept]
    body = src.element.body
    # sectPr (last element) must be preserved
    sectPr = None
    children = list(body.iterchildren())
    if children and children[-1].tag.split("}")[-1] == "sectPr":
        sectPr = children[-1]

    # Re-walk and remove
    p_idx = 0
    pos = 0
    to_remove = []
    for child in list(body.iterchildren()):
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        if pos < first_kept or pos > last_kept:
            to_remove.append(child)
        pos += 1
    for el in to_remove:
        body.remove(el)

    # Insert a cover title at the top
    # Insert paragraphs before the first remaining child
    first_child = next(body.iterchildren(), None)
    cover_paras = []
    for line, size, bold in title_lines:
        p = src.add_paragraph()
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover_paras.append(p._p)
    # add a spacer paragraph
    spacer = src.add_paragraph()._p
    cover_paras.append(spacer)

    # Move cover paragraphs to top (they were appended at end before sectPr)
    for cp in cover_paras:
        body.remove(cp)
    insertion_target = next(body.iterchildren(), None)
    for cp in cover_paras:
        if insertion_target is not None:
            insertion_target.addprevious(cp)
        else:
            body.append(cp)

    src.save(out_path)
    print("wrote", out_path)


# Criteria 1: Business Integration
# Covers: Business Model Development -> Feedback & Analysis (paras 226..273)
build(
    "SENTINEL_Criteria1_Business_Integration.docx",
    [
        ("STA309IAE — Design Thinking and Innovation", 16, True),
        ("SENTINEL Coursework", 13, False),
        ("", 10, False),
        ("Criteria 1: Business Integration", 22, True),
        ("Business Model Development, BPM/BPMN, Product Life Cycle,", 11, False),
        ("Customer Value Chain, and Stakeholder Feedback", 11, False),
        ("", 10, False),
    ],
    keep_p_range=(226, 273),
)

# Criteria 4: Reflection on Outcomes
# Covers: Learning Outcomes (Part A) + Innovation in Business Process (Part B)
#         + Conclusion (Part C) (paras 274..285)
build(
    "SENTINEL_Criteria4_Reflection.docx",
    [
        ("STA309IAE — Design Thinking and Innovation", 16, True),
        ("SENTINEL Coursework", 13, False),
        ("", 10, False),
        ("Criteria 4: Reflection on Outcomes", 22, True),
        ("Part A — Learning Outcomes  |  Part B — Innovation in Business Process", 11, False),
        ("Part C — Conclusion", 11, False),
        ("", 10, False),
    ],
    keep_p_range=(274, 285),
)
