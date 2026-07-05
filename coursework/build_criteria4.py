"""Generate Criteria 4.docx for STA309IAE Design Thinking and Innovation.

Includes the prose for sections B and C plus the infographic prompts
inline so the student can hand them straight to Canva / Bing / DALL-E.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_prompt_box(doc, title, prompt_text):
    """Insert a visually distinct 'Infographic prompt' block."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run(f"🎨  {title}")
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    body = doc.add_paragraph()
    body.paragraph_format.left_indent = Pt(18)
    body_run = body.add_run(prompt_text)
    body_run.italic = True
    body_run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    body.paragraph_format.space_after = Pt(12)


doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# ---- title ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("STA309IAE — Design Thinking and Innovation")
run.bold = True
run.font.size = Pt(16)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = subtitle.add_run("Criteria 4: Reflection on Outcomes")
sub.bold = True
sub.font.size = Pt(13)
sub.font.color.rgb = RGBColor(0x2E, 0x5C, 0xB8)

doc.add_paragraph()

# ---- Section B ----
h1 = doc.add_heading("B. Innovation in Business Process", level=1)
h1.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

doc.add_heading("What is Innovation?", level=2)
doc.add_paragraph(
    "Innovation simply means doing something in a new and better way. It is "
    "not always about inventing a brand new product. Many times innovation "
    "is about changing how a process works inside a company so that the "
    "work becomes faster, cheaper, or gives a better result for the "
    "customer. In business, this kind of change is called business process "
    "innovation."
)

doc.add_heading(
    "Business Process Innovation in the Mobile App Security Sector",
    level=2,
)
doc.add_paragraph(
    "In the mobile application security domain, most companies still "
    "follow a manual penetration testing process. A junior security "
    "researcher has to install the APK, run tools like adb, apktool, jadx "
    "and frida one by one, write notes by hand, and then prepare a long "
    "Word/PDF report at the end. This process takes many days and the "
    "quality depends on how experienced the tester is. Some companies like "
    "NowSecure, Appknox and Quixxi already automate parts of this, but "
    "they mostly focus on static scanning and do not really combine "
    "static, dynamic and AI triage together."
)

doc.add_heading("Why It Was Necessary to Innovate", level=2)
doc.add_paragraph(
    "Our company, SENTINEL, innovates the security testing process by "
    "combining many small steps into one automated pipeline. Instead of "
    "the tester running 10 tools by hand, our platform runs more than 180 "
    "security agents in parallel, uses an LLM to filter false positives, "
    "automatically generates the VAPT report, and even shows step-by-step "
    "reproduction commands for every finding so a junior researcher can "
    "follow easily."
)

doc.add_paragraph(
    "This process innovation helped our organisation in two main ways:"
)

p = doc.add_paragraph(style="List Number")
p.add_run("Value for the customer ").bold = True
p.add_run(
    "— clients get a professional VAPT report in minutes instead of weeks, "
    "with clear evidence, screenshots and reproduction steps."
)

p = doc.add_paragraph(style="List Number")
p.add_run("Revenue generation ").bold = True
p.add_run(
    "— because one scan now takes a few minutes, we can serve many more "
    "clients with the same team. This directly improves our profit margin "
    "and lets us offer a subscription-based business model."
)

add_prompt_box(
    doc,
    "Infographic Prompt 1 — Business Process Innovation Flow",
    "A modern, flat-style horizontal flowchart infographic on a clean "
    "white background, titled 'SENTINEL — Business Process Innovation "
    "Flow'. Six connected hexagonal blocks in a left-to-right pipeline "
    "with arrows between them: (1) APK Upload — mobile phone icon, "
    "(2) Static Agents — magnifying glass over code, (3) Dynamic Agents "
    "— running gear, (4) AI Triage — brain / neural network, (5) Chain "
    "Detection — connected chain links, (6) VAPT Report — document with "
    "shield. Below the pipeline, two large rounded boxes labeled "
    "'Create Value for Customer' (green) and 'Generate Revenue' (blue), "
    "with thin lines connecting from every pipeline block to both boxes. "
    "Use a navy-blue, teal and orange palette, soft drop shadows, "
    "sans-serif typography, minimal text, professional business-report "
    "style.",
)

doc.add_paragraph()

# ---- Section C ----
h2 = doc.add_heading("C. Learning Outcomes", level=1)
h2.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

doc.add_paragraph(
    "From this module I learned many useful things that I did not know "
    "before. First of all, I understood what Design Thinking really means. "
    "Before this module I used to think design is only about how a product "
    "looks, but now I understand it is actually a way of solving problems "
    "by first understanding the user properly. The five steps — Empathise, "
    "Define, Ideate, Prototype, Test — helped me a lot while building my "
    "Sentinel project because I was forced to first think about the "
    "actual user (a junior security researcher) instead of jumping "
    "directly into coding."
)

doc.add_paragraph(
    "I also learned how to think like a business person, not only as a "
    "developer. Things like the Value Proposition Canvas, Business "
    "Process Diagram, Value Chain and Product Life Cycle were completely "
    "new for me. Drawing these diagrams made me realise that a good "
    "technical product can still fail if the business side is weak."
)

doc.add_paragraph(
    "My personal experience while doing this project was that prototyping "
    "and testing again and again is very important. Many of my first "
    "ideas looked good on paper but failed when I tested them. By "
    "following design thinking and updating the prototype every time, the "
    "final solution became much better. Overall, this module taught me "
    "how to combine technology, user empathy and business sense in one "
    "project."
)

add_prompt_box(
    doc,
    "Infographic Prompt 2 — Learning Outcomes Wheel",
    "A circular 'learning outcomes wheel' infographic, flat modern style, "
    "white background. A central circle labeled 'What I Learned in "
    "STA309IAE' with a graduation cap icon. Six pie-slice segments "
    "radiating outward, each in a different pastel color, each containing "
    "an icon and a short label: (1) Design Thinking — lightbulb icon, "
    "(2) User Empathy — heart-with-people icon, (3) Prototyping — "
    "wireframe sketch icon, (4) Business Modeling — chart icon, "
    "(5) Reflection — mirror icon, (6) Teamwork — handshake icon. Around "
    "the outer ring, small dotted lines suggesting iteration. Friendly "
    "student-portfolio aesthetic, rounded sans-serif font, soft "
    "gradients, no clutter.",
)

out = "/home/shambhu/Desktop/project/sentinel/coursework/Criteria_4.docx"
doc.save(out)
print(out)
