"""Convert SENTINEL_THESIS.md to a properly formatted DOCX file."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

def set_cell_background(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def setup_styles(doc):
    # --- Normal ---
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    # --- Heading 1 ---
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Calibri'
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(8)

    # --- Heading 2 ---
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Calibri'
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    # --- Heading 3 ---
    h3 = doc.styles['Heading 3']
    h3.font.name = 'Calibri'
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)

    # --- Code style ---
    if 'Code' not in [s.name for s in doc.styles]:
        code_style = doc.styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = doc.styles['Code']
    code_style.font.name = 'Courier New'
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(4)
    code_style.paragraph_format.left_indent = Inches(0.3)

    # --- Quote style ---
    if 'MyQuote' not in [s.name for s in doc.styles]:
        q_style = doc.styles.add_style('MyQuote', WD_STYLE_TYPE.PARAGRAPH)
    else:
        q_style = doc.styles['MyQuote']
    q_style.font.name = 'Calibri'
    q_style.font.size = Pt(11)
    q_style.font.italic = True
    q_style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    q_style.paragraph_format.left_indent = Inches(0.5)
    q_style.paragraph_format.space_before = Pt(6)
    q_style.paragraph_format.space_after = Pt(6)

def add_cover_page(doc):
    # University name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Softwarica College of IT & E-Commerce')
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run('In Academic Partnership with · Coventry University, United Kingdom')
    run2.font.name = 'Calibri'
    run2.font.size = Pt(12)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

    doc.add_paragraph()
    add_horizontal_rule(doc)
    doc.add_paragraph()

    # Main title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run('SENTINEL')
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(36)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run(
        'Design and Development of an Intelligent Multi-Agent Automated\n'
        'Android Application Security Assessment Framework\n'
        'Using Retrieval-Augmented Generation (RAG) and Large Language Model (LLM)'
    )
    subtitle_run.font.name = 'Calibri'
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.bold = False
    subtitle_run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

    doc.add_paragraph()
    add_horizontal_rule(doc)
    doc.add_paragraph()

    # Details table
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    details = [
        ('Programme', 'BSc (Hons) Ethical Hacking & Cybersecurity'),
        ('Institution', 'Softwarica College × Coventry University UK'),
        ('Method', 'Agile Development with Multi-Agent AI Pipeline'),
        ('Module', 'BSc (Hons) Individual Project · Final Year Research Report'),
        ('Framework', 'SENTINEL — Android APK Security Assessment'),
    ]
    for i, (label, value) in enumerate(details):
        row = table.rows[i]
        set_cell_background(row.cells[0], '1F4E79')
        set_cell_background(row.cells[1], 'DEEAF1')
        lc = row.cells[0].paragraphs[0]
        lr = lc.add_run(label)
        lr.font.bold = True
        lr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        lr.font.name = 'Calibri'
        lr.font.size = Pt(11)
        vc = row.cells[1].paragraphs[0]
        vr = vc.add_run(value)
        vr.font.name = 'Calibri'
        vr.font.size = Pt(11)

    doc.add_paragraph()
    doc.add_paragraph()

    # Stats bar
    stats_p = doc.add_paragraph()
    stats_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stats_run = stats_p.add_run(
        '199 Detection Agents  ·  14 Vulnerability Categories  ·  10-Phase Pipeline\n'
        'RQ1: 79% Coverage Improvement · 59% False Positive Reduction · 8.3min Assessment\n'
        'RQ2: Authorisation Gates · Privacy Controls · LLM Transparency · Human Oversight'
    )
    stats_run.font.name = 'Calibri'
    stats_run.font.size = Pt(10)
    stats_run.font.italic = True
    stats_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    doc.add_page_break()


def parse_and_write(doc, md_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []
    skip_cover = True  # skip first few lines (already handled by cover)
    line_idx = 0

    # Skip the first title block (already on cover page)
    while line_idx < len(lines) and not lines[line_idx].startswith('## Table of Contents'):
        line_idx += 1

    while line_idx < len(lines):
        line = lines[line_idx].rstrip('\n')

        # --- CODE BLOCK ---
        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                # End code block — write it
                in_code_block = False
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph(style='Code')
                # light gray shading
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'F2F2F2')
                pPr.append(shd)
                p.add_run(code_text)
            line_idx += 1
            continue

        if in_code_block:
            code_lines.append(line)
            line_idx += 1
            continue

        # --- TABLE ---
        if line.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            table_rows.append(cells)
            line_idx += 1
            continue
        else:
            if in_table:
                # Render table
                in_table = False
                # Filter out separator rows
                data_rows = [r for r in table_rows if not all(set(c).issubset({'-', ':', ' '}) for c in r)]
                if data_rows:
                    ncols = max(len(r) for r in data_rows)
                    t = doc.add_table(rows=len(data_rows), cols=ncols)
                    t.style = 'Table Grid'
                    for ri, row in enumerate(data_rows):
                        for ci in range(ncols):
                            cell_text = row[ci] if ci < len(row) else ''
                            # Strip markdown bold
                            cell_text = re.sub(r'\*\*(.+?)\*\*', r'\1', cell_text)
                            cell_text = re.sub(r'\*(.+?)\*', r'\1', cell_text)
                            cell_text = re.sub(r'`(.+?)`', r'\1', cell_text)
                            cp = t.rows[ri].cells[ci].paragraphs[0]
                            if ri == 0:
                                set_cell_background(t.rows[ri].cells[ci], '1F4E79')
                                cr = cp.add_run(cell_text)
                                cr.font.bold = True
                                cr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                cr.font.name = 'Calibri'
                                cr.font.size = Pt(10)
                            else:
                                if ri % 2 == 0:
                                    set_cell_background(t.rows[ri].cells[ci], 'DEEAF1')
                                cr = cp.add_run(cell_text)
                                cr.font.name = 'Calibri'
                                cr.font.size = Pt(10)
                    doc.add_paragraph()
                continue

        # --- HORIZONTAL RULE ---
        if line.strip() == '---':
            add_horizontal_rule(doc)
            line_idx += 1
            continue

        # --- PAGE BREAK marker (not in md but skip blank lines after ---)
        if line.strip() == '':
            if line_idx + 1 < len(lines) and lines[line_idx + 1].strip().startswith('## '):
                doc.add_paragraph()
            line_idx += 1
            continue

        # --- HEADINGS ---
        h_match = re.match(r'^(#{1,4})\s+(.*)', line)
        if h_match:
            level = len(h_match.group(1))
            text = h_match.group(2).strip()
            # Clean markdown from heading text
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            if level == 1:
                p = doc.add_heading(text, level=1)
            elif level == 2:
                p = doc.add_heading(text, level=2)
            elif level == 3:
                p = doc.add_heading(text, level=3)
            else:
                p = doc.add_heading(text, level=4) if level == 4 else doc.add_paragraph(text)
            line_idx += 1
            continue

        # --- BLOCKQUOTE ---
        if line.startswith('> '):
            text = line[2:].strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            p = doc.add_paragraph(style='MyQuote')
            p.add_run(text)
            line_idx += 1
            continue

        # --- BULLET LIST ---
        bullet_match = re.match(r'^(\s*[-*+])\s+(.*)', line)
        if bullet_match:
            text = bullet_match.group(2).strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            indent = len(bullet_match.group(1)) - 1
            p = doc.add_paragraph(style='List Bullet' if indent == 0 else 'List Bullet 2')
            p.add_run(text)
            line_idx += 1
            continue

        # --- NUMBERED LIST ---
        num_match = re.match(r'^\d+\.\s+(.*)', line)
        if num_match:
            text = num_match.group(1).strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            p = doc.add_paragraph(style='List Number')
            p.add_run(text)
            line_idx += 1
            continue

        # --- NORMAL PARAGRAPH ---
        text = line.strip()
        if not text:
            line_idx += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

        # Parse inline formatting
        segments = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
        for seg in segments:
            if seg.startswith('**') and seg.endswith('**'):
                r = p.add_run(seg[2:-2])
                r.font.bold = True
                r.font.name = 'Calibri'
                r.font.size = Pt(11)
            elif seg.startswith('*') and seg.endswith('*') and len(seg) > 2:
                r = p.add_run(seg[1:-1])
                r.font.italic = True
                r.font.name = 'Calibri'
                r.font.size = Pt(11)
            elif seg.startswith('`') and seg.endswith('`') and len(seg) > 2:
                r = p.add_run(seg[1:-1])
                r.font.name = 'Courier New'
                r.font.size = Pt(10)
            else:
                # Remove any leftover markdown links [text](url) -> text
                seg = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', seg)
                r = p.add_run(seg)
                r.font.name = 'Calibri'
                r.font.size = Pt(11)

        line_idx += 1


def main():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)

    setup_styles(doc)
    add_cover_page(doc)
    parse_and_write(doc, '/home/shambhu/Desktop/project/sentinel/SENTINEL_THESIS.md')

    out_path = '/home/shambhu/Desktop/project/sentinel/SENTINEL_THESIS.docx'
    doc.save(out_path)
    print(f'Saved: {out_path}')


if __name__ == '__main__':
    main()
