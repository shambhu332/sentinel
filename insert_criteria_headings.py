"""Insert CRITERIA 1 and CRITERIA 4 section headings into SENTINEL_Coursework.docx
to match the split files."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy

SRC = "SENTINEL_Coursework.docx"
doc = Document(SRC)
body = doc.element.body

# Find paragraphs by text prefix so insertion is robust against earlier edits.
TARGETS = [
    {
        "anchor_prefix": "This section describes a conceptual prototype of the SENTINEL platform",
        "occurrence": 1,
        "lines": [
            ("CRITERIA 1", 20, True, True),
            ("BUSINESS INTEGRATION", 18, True, True),
            ("Business Model, Feasibility, BPM/BPMN, Product Life Cycle, Customer Value Chain, and Stakeholder Feedback", 11, False, True),
        ],
    },
    {
        # "Figure 20: Learning Outcomes" appears twice — first in the List of
        # Figures (TOC) and second as the body caption. Anchor on the body one.
        "anchor_prefix": "Figure 20: Learning Outcomes",
        "occurrence": 2,
        "lines": [
            ("CRITERIA 4", 20, True, True),
            ("REFLECTION ON OUTCOMES", 18, True, True),
            ("Part A — Learning Outcomes  |  Part B — Innovation in Business Process  |  Part C — Conclusion", 11, False, True),
        ],
    },
]


def find_paragraph(prefix, occurrence=1):
    """Return the Nth (1-indexed) paragraph whose text starts with prefix."""
    seen = 0
    for child in body.iterchildren():
        if child.tag.split("}")[-1] != "p":
            continue
        text = "".join(child.itertext()).strip()
        if text.startswith(prefix):
            seen += 1
            if seen == occurrence:
                return child
    return None


def make_heading_para(text, size, bold, center):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p._p


for t in TARGETS:
    anchor = find_paragraph(t["anchor_prefix"], t.get("occurrence", 1))
    if anchor is None:
        print("WARN: anchor not found for", t["anchor_prefix"][:50])
        continue
    new_paras = [make_heading_para(*ln) for ln in t["lines"]]
    # add spacer
    new_paras.append(doc.add_paragraph()._p)
    # Move them out of body tail and insert before anchor
    for np in new_paras:
        body.remove(np)
    for np in new_paras:
        anchor.addprevious(np)
    print("inserted before:", t["anchor_prefix"][:50])

doc.save(SRC)
print("updated", SRC)
